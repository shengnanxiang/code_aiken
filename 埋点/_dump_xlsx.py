import sys
import openpyxl

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

path = r"c:\Users\Administrator\Documents\GitHub\code_aiken\埋点\埋点需求-复古相机.xlsx"
wb = openpyxl.load_workbook(path, data_only=True)

print("SHEETS:", wb.sheetnames)

for ws in wb.worksheets:
    print("=" * 100)
    print(f"SHEET: {ws.title}  dims={ws.dimensions}  max_row={ws.max_row}  max_col={ws.max_column}")
    print(f"merged count: {len(ws.merged_cells.ranges)}")
    for r in list(ws.merged_cells.ranges)[:25]:
        print("   merged:", r)
    print("-" * 100)
    limit = min(ws.max_row, 40)
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=limit), start=1):
        cells = []
        for c in row:
            v = c.value
            if v is None:
                continue
            v = str(v).replace("\n", " / ").strip()
            if len(v) > 90:
                v = v[:90] + "..."
            cells.append(f"{c.coordinate}={v}")
        if cells:
            print(f"R{i}: " + " | ".join(cells))
