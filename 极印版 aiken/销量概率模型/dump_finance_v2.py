# -*- coding: utf-8 -*-
"""重新读取财务表 261027，dump Budget 和 Scenario 全部内容找变化"""
import sys
sys.stdout.reconfigure(encoding="utf-8")
import openpyxl
import json

wb = openpyxl.load_workbook(r"C:\Users\Nan\Documents\GitHub\code_aiken\极印版 aiken\AiKen JY Finance_261027.xlsx", data_only=True)
out = {"sheets": wb.sheetnames}

for ws in wb.worksheets:
    rows = []
    for row in ws.iter_rows(values_only=True):
        if row and any(v is not None for v in row):
            rows.append([v for v in row])
    out[ws.title] = rows

with open(r"C:\Users\Nan\Documents\GitHub\code_aiken\极印版 aiken\销量概率模型\_finance_261027.json", "w", encoding="utf-8") as f:
    json.dump(out, f, ensure_ascii=False, indent=1, default=str)
print("sheets:", wb.sheetnames)
for ws in wb.worksheets:
    print(f"  {ws.title}: {ws.max_row} rows")
