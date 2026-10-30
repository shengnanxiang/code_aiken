# -*- coding: utf-8 -*-
"""
更新《AiKen四图直方图数据_PPT用》→ ¥699 口径版
=============================================
与上一版（v0.4 · 599）的差异：
  - 首年/三年两图：数据源换成 s699_v04_results.npz（699 定价 · 漏斗不打折 · v0.4 参数）
  - 参考线锚：按 699 分布的新分位更新（含 P20/P80 备选锚）
  - S1 回测 / DJI 回测两图：不变（与定价无关）
财务注释同步为修正口径：单台净贡献 102.25 元（503.22×70%−250）
"""
import sys
sys.stdout.reconfigure(encoding="utf-8")
import numpy as np
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

D = np.load(r"C:\Users\Nan\Documents\GitHub\code_aiken\极印版 aiken\销量概率模型\s699_v04_results.npz")
y1, y3 = D["y1"], D["y3"]
B = np.load(r"C:\Users\Nan\Documents\GitHub\code_aiken\极印版 aiken\销量概率模型\s1_backtest.npz")
s1_y1 = B["y1"]
rng_d = np.random.default_rng(7)
m_dji = rng_d.lognormal(mean=np.log(600_000), sigma=1.15, size=200_000)

pct = lambda a, q_: float(np.percentile(a, q_))
p10, p20, p50, p80, p90 = pct(y1,10), pct(y1,20), pct(y1,50), pct(y1,80), pct(y1,90)
t10, t50, t90 = pct(y3,10), pct(y3,50), pct(y3,90)

charts = [
    {"sheet": "首年销量分布", "data": y1, "vmax": 45_000, "nbins": 46,
     "ref": (f"参考线：P10={p10:,.0f} / P50={p50:,.0f} / P90={p90:,.0f}（备选锚 P20={p20:,.0f} / P80={p80:,.0f}）"
             f"；Best 阴影从 {p90:,.0f} 起 ｜ ¥699 · 漏斗=v0.4 S1 实测原样（不打折）")},
    {"sheet": "三年累计分布", "data": y3, "vmax": 140_000, "nbins": 46,
     "ref": (f"参考线：P10={t10:,.0f} / P50={t50:,.0f} / P90={t90:,.0f}；10万 阴影从 100,000 起"
             f" ｜ ¥699 口径 ｜ 财务注释：单台净贡献 102.25 元 = 503.22×70% − 250")},
    {"sheet": "S1回测分布", "data": s1_y1, "vmax": 30_000, "nbins": 56,
     "ref": "参考线：内部预测 25,000（P92）/ 实际宽口径 7,285（P5 以下）；阴影到 7,285 止（不变，与定价无关）"},
    {"sheet": "DJI回测分布", "data": m_dji, "vmax": 12_000_000, "nbins": 70,
     "ref": "参考线：内部预期 30-40万 / 上调 100万 / 实售 500万 / 累计 1000万；右尾阴影从 1000万 起（不变，与定价无关）"},
]

wb = openpyxl.Workbook()
wb.remove(wb.active)

HDR_FILL = PatternFill("solid", fgColor="111111")
HDR_FONT = Font(color="FFFFFF", bold=True, size=10)
thin = Side(style="thin", color="CCCCCC")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

for cfg in charts:
    data, vmax, nbins = cfg["data"], cfg["vmax"], cfg["nbins"]
    counts, edges = np.histogram(data, bins=nbins, range=(0, vmax))
    bin_width = vmax / nbins

    # ===== Sheet A：原始数据列（PPT 直方图直接吃） =====
    ws = wb.create_sheet(cfg["sheet"])
    ws.cell(1, 1, f"{cfg['sheet']} · 原始数据列（PPT 直方图用）· ¥699 版").font = Font(bold=True, size=13)
    ws.cell(2, 1, "步骤：① 选 A7 到列尾整列复制 ② PPT 插入→图表→直方图 ③ 粘贴覆盖示例数据").font = Font(size=10, color="666666")
    ws.cell(3, 1, f"④ 右键模轴→设置坐标轴格式→箱宽填 {bin_width:,.0f}（与本国分箱对齐）、㈠ 参考线用手动线条添加").font = Font(size=10, color="EC008C")
    ws.cell(4, 1, cfg["ref"]).font = Font(size=10, bold=True)
    ws.cell(5, 1, f"数据点 {int(counts.sum()):,} 个 = {nbins} 箱 × 频数展开（组中值代表区间）").font = Font(size=9, color="9a9a9a")
    ws.cell(6, 1, "销量(台)").fill = HDR_FILL
    ws.cell(6, 1).font = HDR_FONT

    r = 7
    for i, c in enumerate(counts):
        mid = int(round((edges[i] + edges[i+1]) / 2))
        for _ in range(int(c)):
            ws.cell(r, 1, mid)
            r += 1
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 30
    print(f"[{cfg['sheet']}] 原始数据 {r-7:,} 行")

    # ===== Sheet B：分箱选数（普通柱状图用 · 备用方案） =====
    ws2 = wb.create_sheet(cfg["sheet"] + "_分箱选数")
    ws2.cell(1, 1, f"{cfg['sheet']} · 分箱选数（普通柱状图用 · 备用方案）").font = Font(bold=True, size=13)
    ws2.cell(2, 1, "步骤：① 选中 A5:B 起两列数据 ② PPT 插入→簇状柱状图 ③ 粘贴 ④ 选中柱→间隙宽度调 0% ⑤ X 轴标签会显示区间下限").font = Font(size=10, color="666666")
    ws2.cell(3, 1, "两法画出的图一致；本表优点是数据量小（几十行），缺点是要调间隙宽度。").font = Font(size=10, color="666666")
    ws2.cell(4, 1, cfg["ref"]).font = Font(size=10, bold=True)
    for col, h in [(1, "区间下限"), (2, "模拟次数")]:
        ws2.cell(5, col, h).fill = HDR_FILL
        ws2.cell(5, col, h).font = HDR_FONT
        ws2.cell(5, col, h).border = BORDER
    r2 = 6
    for i, c in enumerate(counts):
        ws2.cell(r2, 1, int(edges[i])).border = BORDER
        ws2.cell(r2, 2, int(c)).border = BORDER
        r2 += 1
    ws2.column_dimensions["A"].width = 12
    ws2.column_dimensions["B"].width = 12

OUT = r"C:\Users\Nan\Documents\GitHub\code_aiken\极印版 aiken\销量概率模型\AiKen四图直方图数据_PPT用.xlsx"
wb.save(OUT)
print("\nwritten:", OUT)
print(f"\n699 锚位速查：首年 P10={p10:,.0f} P20={p20:,.0f} P50={p50:,.0f} P80={p80:,.0f} P90={p90:,.0f}")
print(f"              三年 P10={t10:,.0f} P50={t50:,.0f} P90={t90:,.0f}")
