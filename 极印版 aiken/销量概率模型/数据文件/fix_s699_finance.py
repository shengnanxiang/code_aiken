# -*- coding: utf-8 -*-
"""修复 699 读数表·财务读数 sheet：补扣 30% 销售费用"""
import sys
sys.stdout.reconfigure(encoding="utf-8")
import numpy as np
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

XLSX = r"C:\Users\Nan\Documents\GitHub\code_aiken\极印版 aiken\销量概率模型\AiKen699销量概率分布_读数表.xlsx"
wb = load_workbook(XLSX)

# 删除旧 sheet 重建
if "财务读数" in wb.sheetnames:
    del wb["财务读数"]
ws2 = wb.create_sheet("财务读数", index=1)

d = np.load(r"C:\Users\Nan\Documents\GitHub\code_aiken\极印版 aiken\销量概率模型\s699_v04_results.npz")
y1, y3 = d["y1"], d["y3"]
pct = lambda a, q_: float(np.percentile(a, q_))
prob = lambda a, th: float((a >= th).mean())

# 正确口径
ASP_EX = 503.2181          # 不含税、扣渠道后 ASP（JY 表 BOM&ASP!D10）
SALES_RATE = 0.30          # 销售费用率（用户口径）
BOM = 250.0
UNIT = ASP_EX * (1 - SALES_RATE) - BOM   # 102.25 元/台

HDR_FILL = PatternFill("solid", fgColor="111111")
HDR_FONT = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
B_FONT = Font(name="微软雅黑", size=10, bold=True)
N_FONT = Font(name="微软雅黑", size=10)
S_FONT = Font(name="微软雅黑", size=9, color="6B6B6B")
TITLE_FONT = Font(name="微软雅黑", size=13, bold=True)
THIN = Border(*[Side(style="thin", color="D9D9D9")] * 4)
CENTER = Alignment(horizontal="center", vertical="center")

def style_row(ws, row, cols, font=None, fill=None):
    for c in cols:
        cell = ws.cell(row=row, column=c)
        cell.font = font or N_FONT
        cell.border = THIN
        if fill: cell.fill = fill
        if c > 1: cell.alignment = CENTER

ws2["A1"] = "AiKen ¥699 财务读数（修正版）"
ws2["A1"].font = TITLE_FONT
ws2["A2"] = (f"单台净贡献 = ASP不含税 {ASP_EX:.2f} × (1−销售费用率30%) − BOM {BOM:.0f} = {UNIT:.2f} 元"
             " ｜ ASP 链条：零售 699 → 含税 618.58 → 扣渠道 568.64 → 不含税 503.22（JY 表 BOM&ASP 口径）")
ws2["A2"].font = S_FONT

r = 4
for j, h in enumerate(["指标", "Worst (P10)", "Normal (P50)", "Best (P90)"], 1):
    ws2.cell(row=r, column=j, value=h)
style_row(ws2, r, range(1, 5), HDR_FONT, HDR_FILL)
r += 1
v10, v50, v90 = pct(y1, 10), pct(y1, 50), pct(y1, 90)
t10, t50, t90 = pct(y3, 10), pct(y3, 50), pct(y3, 90)
rows_fin = [
    ("单台净贡献（元）", round(UNIT, 2), round(UNIT, 2), round(UNIT, 2)),
    ("首年销量（台）", round(v10), round(v50), round(v90)),
    ("首年净贡献（万元）", round(v10*UNIT/1e4), round(v50*UNIT/1e4), round(v90*UNIT/1e4)),
    ("三年累计销量（台）", round(t10), round(t50), round(t90)),
    ("三年累计净贡献（万元）", round(t10*UNIT/1e4), round(t50*UNIT/1e4), round(t90*UNIT/1e4)),
]
for row_vals in rows_fin:
    for j, v in enumerate(row_vals, 1):
        ws2.cell(row=r, column=j, value=v)
    style_row(ws2, r, range(1, 5))
    if row_vals[0].startswith("单台"):
        for j in range(1, 5): ws2.cell(row=r, column=j).fill = PatternFill("solid", fgColor="FFF6D9")
    r += 1

r += 1
ws2.cell(row=r, column=1, value="全分布读数").font = B_FONT
r += 1
cm3 = y3 * UNIT
cm1 = y1 * UNIT
NRE = 6_300_000
full_rows = [
    ("P(首年净贡献 ≥ 300万)", f"{prob(cm1,3e6)*100:.1f}%"),
    ("P(首年净贡献 ≥ 500万)", f"{prob(cm1,5e6)*100:.1f}%"),
    ("P(三年累计净贡献 ≥ 0)", f"{prob(cm3,0)*100:.1f}%"),
    ("P(三年累计净贡献 ≥ 300万)", f"{prob(cm3,3e6)*100:.1f}%"),
    ("P(三年累计净贡献 ≥ 500万)", f"{prob(cm3,5e6)*100:.1f}%"),
    ("P(三年累计净贡献 ≥ 1000万)", f"{prob(cm3,1e7)*100:.1f}%"),
    ("P(三年累计 ≥ 10万台)", f"{prob(y3,1e5)*100:.1f}%"),
]
for label, v in full_rows:
    ws2.cell(row=r, column=1, value=label)
    ws2.cell(row=r, column=2, value=v)
    style_row(ws2, r, range(1, 3))
    r += 1

r += 1
ws2.cell(row=r, column=1, value="注 1：简化口径 = 单台净贡献 × 销量（102.25 元/台），已含 30% 销售费用扣除；").font = S_FONT
r += 1
ws2.cell(row=r, column=1, value="        不含物流/MOH/返修/模具摊销/资金成本等细项——精确口径以 JY 699 Finance 表 Scenario sheets 为准").font = S_FONT
r += 1
ws2.cell(row=r, column=1, value="注 2：JY 表完整成本链三档净利（三年）：Normal +78.7万 · Best +392.9万 · Worst −94.5万；").font = S_FONT
r += 1
ws2.cell(row=r, column=1, value="        两口径差异 = JY 表多扣了物流(1.16/unit)/MOH(12.5)/返修(4.25)/服务费(4.25)/资金成本及模具摊销").font = S_FONT

for col, w_ in zip("ABCD", [40, 16, 16, 16]):
    ws2.column_dimensions[col].width = w_

wb.save(XLSX)
print("财务读数 sheet 已重建（修正版）")
print(f"\n修正后关键数字：")
print(f"单台净贡献：{UNIT:.2f} 元（此前错误值 253.22）")
print(f"首年净贡献：P10={v10*UNIT/1e4:.0f}万  P50={v50*UNIT/1e4:.0f}万  P90={v90*UNIT/1e4:.0f}万")
print(f"三年累计净贡献：P10={t10*UNIT/1e4:.0f}万  P50={t50*UNIT/1e4:.0f}万  P90={t90*UNIT/1e4:.0f}万")
print(f"P(三年累计净贡献≥0) = {prob(cm3,0)*100:.1f}%")
print(f"P(三年累计净贡献≥500万) = {prob(cm3,5e6)*100:.1f}%")
