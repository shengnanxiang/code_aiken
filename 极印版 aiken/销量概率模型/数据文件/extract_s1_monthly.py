# -*- coding: utf-8 -*-
"""提取 S1 的月度投放与转化序列（JD 站内明细表）"""
import sys
sys.stdout.reconfigure(encoding="utf-8")
from openpyxl import load_workbook
import collections

wb = load_workbook(r"C:\Users\Nan\Documents\GitHub\code_aiken\极印版 aiken\销量概率模型\JD26年效果广告分sku_站内.xlsx", data_only=True)
ws = wb["明细表"]

# 列位置：日期(1) 型号(4) 展现(5) 点击(6) 花费(8) 直接订单行数(10) 直接加购(13)
rows = ws.iter_rows(min_row=2, values_only=True)
monthly = collections.defaultdict(lambda: dict(exp=0, clk=0, cost=0.0, orders=0, carts=0))
skus = collections.Counter()
for r in rows:
    if r is None or len(r) < 14: continue
    date, model = r[0], r[3]
    if not isinstance(date, str) or model not in ("S1",): 
        if isinstance(model, str): skus[model] += 1
        continue
    # 日期形如 20260301~20260331
    month = date[:6]
    d = monthly[month]
    d["exp"] += r[4] or 0
    d["clk"] += r[5] or 0
    d["cost"] += r[7] or 0
    d["orders"] += r[9] or 0
    d["carts"] += r[12] if len(r) > 12 and isinstance(r[12], (int, float)) else 0

print("S1 各月（JD 站内直接归因）：")
print(f"{'月份':<8}{'展现':>12}{'点击':>10}{'CTR':>8}{'花费':>12}{'订单':>8}{'CVR':>9}{'加购':>8}")
tot = dict(exp=0, clk=0, cost=0.0, orders=0, carts=0)
for month in sorted(monthly):
    d = monthly[month]
    ctr = d['clk']/d['exp']*100 if d['exp'] else 0
    cvr = d['orders']/d['clk']*100 if d['clk'] else 0
    print(f"{month:<8}{d['exp']:>12,}{d['clk']:>10,}{ctr:>7.2f}%{d['cost']:>12,.0f}{d['orders']:>8}{cvr:>8.2f}%{d['carts']:>8}")
    for k in tot: tot[k] += d[k]
print(f"{'合计':<8}{tot['exp']:>12,}{tot['clk']:>10,}{tot['clk']/tot['exp']*100:>7.2f}%{tot['cost']:>12,.0f}{tot['orders']:>8}{tot['orders']/tot['clk']*100:>8.2f}%{tot['carts']:>8}")
print(f"\nCPM(加权) = {tot['cost']/tot['exp']*1000:.1f} 元")
print(f"数据覆盖月份: {sorted(monthly)}")
