# -*- coding: utf-8 -*-
"""验证 S1 全渠道销量的构成：JD间接订单 + 站外投放转化换算"""
import sys
sys.stdout.reconfigure(encoding="utf-8")
from openpyxl import load_workbook
import collections

# ===== 1. JD 站内：直接 vs 间接订单 =====
wb = load_workbook(r"C:\Users\Nan\Documents\GitHub\code_aiken\极印版 aiken\销量概率模型\JD26年效果广告分sku_站内.xlsx", data_only=True)
ws = wb["明细表"]

monthly = collections.defaultdict(lambda: dict(direct=0, indirect=0, dgmv=0.0, igmv=0.0))
for r in ws.iter_rows(min_row=2, values_only=True):
    if r is None or len(r) < 15: continue
    date, model = r[0], r[3]
    if not isinstance(date, str) or model != "S1": continue
    month = date[:6]
    d = monthly[month]
    d["direct"] += r[9] or 0          # 直接订单行数
    d["indirect"] += r[11] or 0       # 间接订单行数
    d["dgmv"] += r[10] or 0           # 直接订单金额
    d["igmv"] += r[12] or 0           # 间接订单金额

tot_d = sum(d["direct"] for d in monthly.values())
tot_i = sum(d["indirect"] for d in monthly.values())
tot_dg = sum(d["dgmv"] for d in monthly.values())
tot_ig = sum(d["igmv"] for d in monthly.values())
print("=== JD 站内（S1，2026-03~10）===")
print(f"直接订单: {tot_d} 台 ｜ GMV {tot_dg:,.0f} 元")
print(f"间接订单: {tot_i} 台 ｜ GMV {tot_ig:,.0f} 元")
print(f"合计: {tot_d + tot_i} 台 ｜ GMV {tot_dg+tot_ig:,.0f} 元")
print(f"间接/直接 = {tot_i/tot_d:.2f}x")

# ===== 2. 站外投放转化估算 =====
# 站外 S1 归属花费（小红书+抖音）= 626,289 + 904,104 = 1,530,393 元
# 用 JD 站内 ROI 做锚：站内 ROI 8.25（GMV/花费）——但站外种草的转化路径更间接
# 三种口径：
OUT_SPEND = 626_289.4 + 904_104.0   # 153.0 万
IN_ROI = 8.254                       # JD 站内 GMV ROI（汇总表）
PRICE = 1799.0
print(f"\n=== 站外投放（S1 归属花费 {OUT_SPEND/1e4:.1f} 万）===")
for label, roi in [("站外 ROI = 站内 8.25（乐观）", IN_ROI),
                   ("站外 ROI = 站内 50%（种草衰减）", IN_ROI * 0.5),
                   ("站外 ROI = 站内 30%（保守）", IN_ROI * 0.3)]:
    gmv = OUT_SPEND * roi
    units = gmv / PRICE
    print(f"{label}: GMV {gmv/1e4:,.0f} 万 → {units:,.0f} 台")

# ===== 3. 汇总：首年外推的几个口径 =====
months_active = 7  # 2026-04~10 有效投放月（3月只有10天测试期？3月展现82万但有费用2.35万）
# 实际上 3-10 是 8 个月，检查每月都有投放
print(f"\n覆盖月份: {sorted(monthly)}（共 {len(monthly)} 个月）")
n_m = len(monthly)
jd_all = tot_d + tot_i
jd_monthly = jd_all / n_m
for out_mult, label in [(0.0, "仅 JD 直接+间接"),
                        (0.3, "JD + 站外(30%效率)"),
                        (0.5, "JD + 站外(50%效率)"),
                        (1.0, "JD + 站外(100%效率)")]:
    out_units = OUT_SPEND * IN_ROI * out_mult / PRICE if out_mult > 0 else 0
    ytd = jd_all + out_units
    y1 = ytd / n_m * 12
    print(f"{label:24s}: YTD {ytd:,.0f} 台 → 首年外推 {y1:,.0f} 台")
