# -*- coding: utf-8 -*-
import sys
sys.stdout.reconfigure(encoding="utf-8")
from openpyxl import load_workbook
import collections

wb = load_workbook(r"C:\Users\Nan\Documents\GitHub\code_aiken\极印版 aiken\销量概率模型\JD26年效果广告分sku_站内.xlsx", data_only=True)
ws = wb["明细表"]

# 硬编码列位（从表头打印确认）：date=0, model=3, direct=10, dgmv=11, indirect=12, igmv=13
monthly = collections.defaultdict(lambda: dict(direct=0, indirect=0, dgmv=0.0, igmv=0.0))
for r in ws.iter_rows(min_row=2, values_only=True):
    if r is None or len(r) < 14: continue
    date, model = r[0], r[3]
    if not isinstance(date, str) or model != "S1": continue
    month = date[:6]
    d = monthly[month]
    d["direct"] += r[10] or 0
    d["indirect"] += r[12] or 0
    d["dgmv"] += r[11] or 0
    d["igmv"] += r[13] or 0

print("=== S1 各月（JD 站内）===")
tot = dict(direct=0, indirect=0, dgmv=0.0, igmv=0.0)
for m in sorted(monthly):
    d = monthly[m]
    print(f"{m}: 直接 {d['direct']:>7.1f}  间接 {d['indirect']:>7.1f}  直接GMV {d['dgmv']:>10,.0f}  间接GMV {d['igmv']:>10,.0f}")
    for k in tot: tot[k] += d[k]
print(f"合计: 直接 {tot['direct']:>7.1f}  间接 {tot['indirect']:>7.1f}  直接GMV {tot['dgmv']:>10,.0f}  间接GMV {tot['igmv']:>10,.0f}")

OUT_SPEND = 626_289.4 + 904_104.0
IN_ROI = 8.254
PRICE = 1799.0
n_m = len(monthly)
jd_all = tot["direct"] + tot["indirect"]
print(f"\n=== 首年外推（覆盖 {n_m} 个月）===")
for out_mult, label in [(0.0, "仅JD站内"), (0.3, "JD+站外30%"), (0.5, "JD+站外50%"), (1.0, "JD+站外100%")]:
    out_units = OUT_SPEND * IN_ROI * out_mult / PRICE if out_mult > 0 else 0
    ytd = jd_all + out_units
    y1 = ytd / n_m * 12
    print(f"{label}: YTD {ytd:,.0f} 台 -> 首年外推 {y1:,.0f} 台")
