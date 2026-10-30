# -*- coding: utf-8 -*-
"""读取 Ricotta 599 IL Finance 表，理解结构与实际销量数据"""
import sys
sys.stdout.reconfigure(encoding="utf-8")
from openpyxl import load_workbook

PATH = r"D:\OneDrive\01 汉图\03 产品\04 Ricotta\IL\Ricotta 599 IL Finance_241206.xlsx"
wb = load_workbook(PATH, data_only=True)
print("Sheets:", wb.sheetnames)
for sn in wb.sheetnames:
    ws = wb[sn]
    print(f"\n===== Sheet: {sn} ({ws.max_row}r x {ws.max_column}c) =====")
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 70)):
        vals = []
        for c in row:
            if c.value is not None:
                v = c.value
                if isinstance(v, float):
                    v = round(v, 4)
                vals.append(f"{c.coordinate}:{v}")
        if vals:
            print(" | ".join(vals)[:400])
