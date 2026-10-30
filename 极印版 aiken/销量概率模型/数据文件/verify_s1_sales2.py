# -*- coding: utf-8 -*-
"""验证 S1 全渠道销量构成（修正列位）——先打印表头列位再累计"""
import sys
sys.stdout.reconfigure(encoding="utf-8")
from openpyxl import load_workbook
import collections

wb = load_workbook(r"C:\Users\Nan\Documents\GitHub\code_aiken\极印版 aiken\销量概率模型\JD26年效果广告分sku_站内.xlsx", data_only=True)
ws = wb["明细表"]

# 先打印表头，确认列名
hdr = [c.value for c in ws[1]]
print("表头列位：")
for i, h in enumerate(hdr):
    if h: print(f"  col{i}: {h}")

# 找列
col_map = {}
for i, h in enumerate(hdr):
    if h is None: continue
    hs = str(h)
    if "直接订单行数" in hs and "预" not in hs and "direct" not in col_map: col_map["direct"] = i
    elif "间接订单行数" in hs and "indirect" not in col_map: col_map["indirect"] = i
    elif "直接订单金额" in hs and "dgmv" not in col_map: col_map["dgmv"] = i
    elif "间接订单金额" in hs and "igmv" not in col_map: col_map["igmv"] = i
    elif hs == "型号" and "model" not in col_map: col_map["model"] = i
    elif hs == "日期" and "date" not in col_map: col_map["date"] = i
print(f"\n列映射: {col_map}")

monthly = collections.defaultdict(lambda: dict(direct=0, indirect=0, dgmv=0.0, igmv=0.0))
for r in ws.iter_rows(min_row=2, values_only=True):
    if r is None or len(r) <= max(col_map.values()): continue
    date, model = r[col_map["date"]], r[col_map["model"]]
    if not isinstance(date, str) or model != "S1": continue
    month = date[:6]
    d = monthly[month]
    d["direct"] += r[col_map["direct"]] or 0
    d["indirect"] += r[col_map["indirect"]] or 0
    d["dgmv"] += r[col_map["dgmv"]] or 0
    d["igmv"] += r[col_map["igmv"]] or 0

print("\n=== S1 各月（JD 站内）===")
print(f"{'月份':<8}{'直接':>8}{'间接':>8}{'直接GMV':>12}{'间接GMV':>12}")
tot = dict(direct=0, indirect=0, dgmv=0.0, igmv=0.0)
for m in sorted(monthly):
    d = monthly[m]
    print(f"{m:<8}{d['direct']:>8.1f}{d['indirect']:>8.1f}{d['dgmv']:>12,.0f}{d['igmv']:>12,.0f}")
    for k in tot: tot[k] += d[k]
print(f"{'合计':<8}{tot['direct']:>8.1f}{tot['indirect']:>8.1f}{tot['dgmv']:>12,.0f}{tot['igmv']:>12,.0f}")

# ===== 站外转化与首年外推 =====
OUT_SPEND = 626_289.4 + 904_104.0
IN_ROI = 8.254
PRICE = 1799.0
n_m = len(monthly)
jd_all = tot["direct"] + tot["indirect"]

print(f"\n=== 首年外推（覆盖 {n_m} 个月）===")
for out_mult, label in [(0.0, "仅 JD 站内"),
                        (0.3, "JD + 站外30%"),
                        (0.5, "JD + 站外50%"),
                        (1.0, "JD + 站外100%")]:
    out_units = OUT_SPEND * IN_ROI * out_mult / PRICE if out_mult > 0 else 0
    ytd = jd_all + out_units
    y1 = ytd / n_m * 12
    print(f"{label:14s}: YTD {ytd:,.0f} 台 → 首年外推 {y1:,.0f} 台")
