# -*- coding: utf-8 -*-
"""读取两份投放数据，看 S1 的真实月度表现（找校准用的观测）"""
import sys
sys.stdout.reconfigure(encoding="utf-8")
from openpyxl import load_workbook

for path, label in [
    (r"C:\Users\Nan\Documents\GitHub\code_aiken\极印版 aiken\销量概率模型\JD26年效果广告分sku_站内.xlsx", "JD站内"),
    (r"C:\Users\Nan\Documents\GitHub\code_aiken\极印版 aiken\销量概率模型\极印投放数据_站外2026YTD1026.xlsx", "站外YTD"),
]:
    wb = load_workbook(path, data_only=True)
    print(f"\n{'='*70}\n{label}: {wb.sheetnames}")
    for sn in wb.sheetnames:
        ws = wb[sn]
        print(f"\n--- {sn} ({ws.max_row}r x {ws.max_column}c) ---")
        # 打印前 25 行、每行前 14 列的非空值
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 25), max_col=min(ws.max_column, 14)):
            vals = []
            for c in row:
                if c.value is not None:
                    v = c.value
                    if isinstance(v, float): v = round(v, 3)
                    vals.append(str(v)[:22])
            if vals:
                print(" | ".join(vals))
