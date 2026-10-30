# -*- coding: utf-8 -*-
import sys
sys.stdout.reconfigure(encoding="utf-8")
import openpyxl
wb = openpyxl.load_workbook(r"C:\Users\Nan\Documents\GitHub\code_aiken\极印版 aiken\销量概率模型\AiKen销量概率分布_读数表.xlsx", data_only=True)
print("sheets:", wb.sheetnames)
ws = wb[wb.sheetnames[-1]] if len(wb.sheetnames) > 1 else wb.active
for name in wb.sheetnames:
    ws = wb[name]
    for row in ws.iter_rows(values_only=True):
        if not row: continue
        s0 = str(row[0]) if row[0] is not None else ""
        if s0 in ("P10", "P50", "P90"):
            print(f"[{name}] {s0}: 首年 {row[1]:,} | 三年 {row[4]:,} | {row[7] if len(row)>7 and row[7] else ''}")
        if "v0.4" in s0 or "v0.3" in s0:
            print(f"[{name}] 版本: {s0[:50]}")
