# -*- coding: utf-8 -*-
import openpyxl
wb = openpyxl.load_workbook(r"C:\Users\Nan\Documents\GitHub\code_aiken\极印版 aiken\AiKen JY Finance_261027.xlsx", data_only=True)
ws = wb["Budget"]
out = []
targets = {634700: None, 1822700: None, 4138050: None, 4773550: None, 1188000: None, 1650000: None, 3003000: None}
for row in ws.iter_rows(min_row=1, max_row=40):
    for c in row:
        if c.value in targets:
            out.append(f"FOUND {c.coordinate} = {c.value}")
out.append("--- 右侧汇总区 (E-J, 行1-14) ---")
for r in range(1, 15):
    cells = []
    for col in "EFGHIJ":
        v = ws[f"{col}{r}"].value
        if v is not None:
            cells.append(f"{col}{r}={v}")
    if cells:
        out.append(" | ".join(str(x) for x in cells))
with open(r"C:\Users\Nan\Documents\GitHub\code_aiken\极印版 aiken\销量概率模型\_cell_check.txt", "w", encoding="utf-8") as f:
    f.write("\n".join(out))
print("done", len(out))
