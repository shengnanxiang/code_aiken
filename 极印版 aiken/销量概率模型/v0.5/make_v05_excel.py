# -*- coding: utf-8 -*-
"""v0.5 三件套之一：直方图数据 Excel（PPT 用）+ 之二：读数表"""
import sys
sys.stdout.reconfigure(encoding="utf-8")
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

D = np.load(r"C:\Users\Nan\Documents\GitHub\code_aiken\极印版 aiken\销量概率模型\v0.5\v05_results.npz")
y1, y3 = D["y1"], D["y3"]
B = np.load(r"C:\Users\Nan\Documents\GitHub\code_aiken\极印版 aiken\销量概率模型\s1_backtest.npz")
s1_y1 = B["y1"]
rng_d = np.random.default_rng(7)
m_dji = rng_d.lognormal(mean=np.log(600_000), sigma=1.15, size=200_000)

pct = lambda a, x: float(np.percentile(a, x))
prob = lambda a, th: float((a >= th).mean())
UNIT = 503.2181 * 0.7 - 250

HDR_FILL = PatternFill("solid", fgColor="111111")
ANCHOR_FILL = PatternFill("solid", fgColor="FFF6D9")
HDR_FONT = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
B_FONT = Font(name="微软雅黑", size=10, bold=True)
N_FONT = Font(name="微软雅黑", size=10)
S_FONT = Font(name="微软雅黑", size=9, color="6B6B6B")
TITLE_FONT = Font(name="微软雅黑", size=13, bold=True)
THIN = Border(*[Side(style="thin", color="D9D9D9")] * 4)
CENTER = Alignment(horizontal="center", vertical="center")
WRAP = Alignment(wrap_text=True, vertical="top")

# =========================================================
# 件1：直方图数据 Excel（四图：首年/三年/S1回测/DJI回测）
# =========================================================
wb = Workbook()
wb.remove(wb.active)

charts = [
    {"sheet": "首年销量分布", "data": y1, "vmax": 45_000, "nbins": 46,
     "ref": (f"参考线：Worst锚 P20={pct(y1,20):,.0f} / Normal锚 P50={pct(y1,50):,.0f} / Best锚 P90={pct(y1,90):,.0f}"
             f"；Best 阴影从 {pct(y1,90):,.0f} 起 ｜ v0.5：核心对标带 ¥500-800 校准")},
    {"sheet": "三年累计分布", "data": y3, "vmax": 120_000, "nbins": 46,
     "ref": (f"参考线：P10={pct(y3,10):,.0f} / P50={pct(y3,50):,.0f} / P90={pct(y3,90):,.0f}；10万 阴影从 100,000 起"
             f" ｜ v0.5 ｜ 财务：单台净贡献 102.25 元 = 503.22×70% − 250")},
    {"sheet": "S1回测分布", "data": s1_y1, "vmax": 30_000, "nbins": 56,
     "ref": "参考线：内部预测 25,000（P92）/ 实际宽口径 7,285（P5 以下）；阴影到 7,285 止（不变，与定价无关）"},
    {"sheet": "DJI回测分布", "data": m_dji, "vmax": 12_000_000, "nbins": 70,
     "ref": "参考线：内部预期 30-40万 / 上调 100万 / 实售 500万 / 累计 1000万；右尾阴影从 1000万 起（不变，与定价无关）"},
]

for cfg in charts:
    data, vmax, nbins = cfg["data"], cfg["vmax"], cfg["nbins"]
    counts, edges = np.histogram(data, bins=nbins, range=(0, vmax))
    bin_width = vmax / nbins

    ws = wb.create_sheet(cfg["sheet"])
    ws.cell(1, 1, f"{cfg['sheet']} · 原始数据列（PPT 直方图用）· v0.5").font = TITLE_FONT
    ws.cell(2, 1, "步骤：① 选 A7 到列尾整列复制 ② PPT 插入→图表→直方图 ③ 粘贴覆盖示例数据").font = Font(name="微软雅黑", size=10, color="666666")
    ws.cell(3, 1, f"④ 右键模轴→设置坐标轴格式→箱宽填 {bin_width:,.0f}（与本国分箱对齐）、㈠ 参考线用手动线条添加").font = Font(name="微软雅黑", size=10, color="EC008C")
    ws.cell(4, 1, cfg["ref"]).font = Font(name="微软雅黑", size=10, bold=True)
    ws.cell(5, 1, f"数据点 {int(counts.sum()):,} 个 = {nbins} 箱 × 频数展开（组中值代表区间）").font = S_FONT
    ws.cell(6, 1, "销量(台)").fill = HDR_FILL
    ws.cell(6, 1).font = HDR_FONT
    r = 7
    for i, c in enumerate(counts):
        mid = int(round((edges[i] + edges[i+1]) / 2))
        for _ in range(int(c)):
            ws.cell(r, 1, mid)
            r += 1
    ws.column_dimensions["A"].width = 12
    print(f"[{cfg['sheet']}] {r-7:,} 行")

    ws2 = wb.create_sheet(cfg["sheet"] + "_分箱选数")
    ws2.cell(1, 1, f"{cfg['sheet']} · 分箱选数（普通柱状图用 · 备用）").font = TITLE_FONT
    ws2.cell(2, 1, "步骤：① 选中 A5:B 起两列 ② 插入→簇状柱状图 ③ 粘贴 ④ 间隙宽度调 0%").font = Font(name="微软雅黑", size=10, color="666666")
    ws2.cell(4, 1, cfg["ref"]).font = Font(name="微软雅黑", size=10, bold=True)
    for col, h in [(1, "区间下限"), (2, "模拟次数")]:
        ws2.cell(5, col, h).fill = HDR_FILL
        ws2.cell(5, col, h).font = HDR_FONT
        ws2.cell(5, col, h).border = THIN
    r2 = 6
    for i, c in enumerate(counts):
        ws2.cell(r2, 1, int(edges[i])).border = THIN
        ws2.cell(r2, 2, int(c)).border = THIN
        r2 += 1
    ws2.column_dimensions["A"].width = 12
    ws2.column_dimensions["B"].width = 12

OUT1 = r"C:\Users\Nan\Documents\GitHub\code_aiken\极印版 aiken\销量概率模型\v0.5\AiKen四图直方图数据_PPT用_v05.xlsx"
wb.save(OUT1)
print("saved:", OUT1)

# =========================================================
# 件2：读数表（v0.5 口径，锚 Worst=P20 / Normal=P50 / Best=P90）
# =========================================================
wb2 = Workbook()

# ---- Sheet 1: 销量分位表 ----
ws = wb2.active
ws.title = "销量分位表"
ws["B2"] = "AiKen ¥699 销量分位表 · v0.5（核心对标带 ¥500-800 校准）"
ws["B2"].font = TITLE_FONT
ws["B3"] = ("读法：「首年 P50 = X」表示 50% 的模拟结果首年销量低于 X。黄色底 = 三档锚（用户裁决：Worst=P20 / Normal=P50 / Best=P90）。"
            "v0.5 变更：m 先验 11万/σ0.85 → 9.5万/σ0.80（参照带收窄到核心对标带：instax mini 12 主锚 + Flashback 下界 + Paper Shoot 低分位；"
            "Kodak Mini Shot 移出参数仅作观察）。漏斗/财务口径与 v0.4 一致。")
ws["B3"].font = S_FONT

QUANTILES = [5, 10, 15, 20, 25, 30, 35, 40, 45, 50, 55, 60, 65, 70, 75, 80, 85, 90, 95]
r = 5
for j, h in enumerate(["分位 P(x)", "首年销量（台）", "首年（万台）", "首年 ≥ 该销量概率", "三年累计（台）", "三年（万台）", "三年 ≥ 该销量概率", "备注"], 2):
    ws.cell(row=r, column=j, value=h)
style_r = range(2, 10)
for j in style_r:
    ws.cell(row=r, column=j).fill = HDR_FILL
    ws.cell(row=r, column=j).font = HDR_FONT
    ws.cell(row=r, column=j).border = THIN
r += 1
ANCHORS = {20: ("Worst 锚（用户裁决）", 0), 50: ("Normal 锚", 0), 90: ("Best 锚", 0)}
for qn in QUANTILES:
    v1, v3 = pct(y1, qn), pct(y3, qn)
    ws.cell(row=r, column=2, value=f"P{qn}")
    ws.cell(row=r, column=3, value=round(v1))
    ws.cell(row=r, column=4, value=round(v1/1e4, 2))
    ws.cell(row=r, column=5, value=f"{100-qn}%")
    ws.cell(row=r, column=6, value=round(v3))
    ws.cell(row=r, column=7, value=round(v3/1e4, 2))
    ws.cell(row=r, column=8, value=f"{100-qn}%")
    if qn in ANCHORS:
        label = ANCHORS[qn][0]
        ws.cell(row=r, column=9, value=f"{label}：首年平均月销 ~{v1/12:,.0f}（三年月均 ~{v3/36:,.0f}）")
    for j in style_r:
        ws.cell(row=r, column=j).border = THIN
        ws.cell(row=r, column=j).font = N_FONT
        if j > 2: ws.cell(row=r, column=j).alignment = CENTER
    if qn in ANCHORS:
        for j in style_r:
            ws.cell(row=r, column=j).fill = ANCHOR_FILL
            ws.cell(row=r, column=j).font = B_FONT
    r += 1

r += 1
ws.cell(row=r, column=2, value="关键概率读数").font = B_FONT
r += 1
KEYS = [
    ("P(首年 ≥ Worst锚 16,678)", pct(y1,20)),
    ("P(首年 ≥ Normal锚 21,827)", pct(y1,50)),
    ("P(首年 ≥ Best锚 32,307)", pct(y1,90)),
    ("P(首年 ≥ 3.0万)", 30_000),
    ("P(三年累计 ≥ 6万)", 60_000),
    ("P(三年累计 ≥ 10万 · 十万台线)", 100_000),
]
for label, th in KEYS:
    ws.cell(row=r, column=2, value=label)
    which_y1 = "首年" in label
    ws.cell(row=r, column=3, value=f"{prob(y1 if which_y1 else y3, th)*100:.1f}%")
    for j in (2, 3):
        ws.cell(row=r, column=j).border = THIN
    r += 1

for col, w_ in zip("BCDEFGHI", [32, 15, 13, 17, 15, 13, 17, 36]):
    ws.column_dimensions[col].width = w_

# ---- Sheet 2: 财务读数 ----
ws2 = wb2.create_sheet("财务读数")
ws2["A1"] = "AiKen ¥699 财务读数（v0.5）"
ws2["A1"].font = TITLE_FONT
ws2["A2"] = (f"单台净贡献 = ASP不含税 503.22 × (1−销售费用率30%) − BOM 250 = {UNIT:.2f} 元"
             " ｜ ASP 链：零售 699 → 含税 618.58 → 扣渠道 568.64 → 不含税 503.22")
ws2["A2"].font = S_FONT
r = 4
for j, h in enumerate(["指标", "Worst (P20)", "Normal (P50)", "Best (P90)"], 1):
    ws2.cell(row=r, column=j, value=h)
    ws2.cell(row=r, column=j).fill = HDR_FILL
    ws2.cell(row=r, column=j).font = HDR_FONT
    ws2.cell(row=r, column=j).border = THIN
r += 1
vw, vn, vb = pct(y1,20), pct(y1,50), pct(y1,90)
tw, tn, tb = pct(y3,20), pct(y3,50), pct(y3,90)
rows_fin = [
    ("单台净贡献（元）", round(UNIT, 2), round(UNIT, 2), round(UNIT, 2)),
    ("首年销量（台）", round(vw), round(vn), round(vb)),
    ("首年净贡献（万元）", round(vw*UNIT/1e4), round(vn*UNIT/1e4), round(vb*UNIT/1e4)),
    ("三年累计销量（台）", round(tw), round(tn), round(tb)),
    ("三年累计净贡献（万元）", round(tw*UNIT/1e4), round(tn*UNIT/1e4), round(tb*UNIT/1e4)),
]
for row_vals in rows_fin:
    for j, v in enumerate(row_vals, 1):
        ws2.cell(row=r, column=j, value=v)
        ws2.cell(row=r, column=j).border = THIN
        ws2.cell(row=r, column=j).font = N_FONT
        if j > 1: ws2.cell(row=r, column=j).alignment = CENTER
    if row_vals[0].startswith("单台"):
        for j in range(1, 5): ws2.cell(row=r, column=j).fill = ANCHOR_FILL
    r += 1

r += 1
ws2.cell(row=r, column=1, value="全分布读数").font = B_FONT
r += 1
cm3 = y3 * UNIT
cm1 = y1 * UNIT
full_rows = [
    ("P(首年净贡献 ≥ 200万)", f"{prob(cm1,2e6)*100:.1f}%"),
    ("P(首年净贡献 ≥ 300万)", f"{prob(cm1,3e6)*100:.1f}%"),
    ("P(三年累计净贡献 ≥ 0)", f"{prob(cm3,0)*100:.1f}%"),
    ("P(三年累计净贡献 ≥ 300万)", f"{prob(cm3,3e6)*100:.1f}%"),
    ("P(三年累计净贡献 ≥ 500万)", f"{prob(cm3,5e6)*100:.1f}%"),
    ("P(三年累计净贡献 ≥ 1000万)", f"{prob(cm3,1e7)*100:.1f}%"),
]
for label, v in full_rows:
    ws2.cell(row=r, column=1, value=label)
    ws2.cell(row=r, column=2, value=v)
    for j in (1, 2): ws2.cell(row=r, column=j).border = THIN
    r += 1
r += 1
ws2.cell(row=r, column=1, value="注：简化口径 = 单台净贡献 × 销量（102.25 元/台，已含 30% 销售费用扣除）；").font = S_FONT
r += 1
ws2.cell(row=r, column=1, value="        不含物流/MOH/返修/模具摊销/资金成本——精确口径以 JY 699 Finance 表 Scenario sheets 为准").font = S_FONT
for col, w_ in zip("ABCD", [40, 16, 16, 16]):
    ws2.column_dimensions[col].width = w_

# ---- Sheet 3: 分布说明 ----
ws3 = wb2.create_sheet("分布说明")
ws3["B2"] = "AiKen ¥699 销量概率分布 · 读数表（v0.5）"
ws3["B2"].font = TITLE_FONT
ws3["B3"] = "模型 v0.5 · 8,000 次蒙特卡洛 · 核心对标带 ¥500-800 校准 · 锚：Worst=P20 / Normal=P50 / Best=P90"
ws3["B3"].font = S_FONT
notes = [
    ("v0.5 与 v0.4 的差异是什么？", "只动了 m（市场潜力）先验：11万/σ0.85 → 9.5万/σ0.80。参照带从「¥400-1,300 全带」收窄到「核心对标带 ¥500-800」——主锚 instax mini 12（¥599-699，中国年销估 25-40 万），下界 Flashback One35（¥749，三年用户 5-10 万），低分位 Paper Shoot。Kodak Mini Shot（¥1,269）移出参数区，仅作观察带水位证据。q、漏斗（S1 实测）、财务口径均未变。"),
    ("三档锚（v0.5 起的用户裁决）", "Worst = P20（首年跌破该线的概率 20%，即 80% 置信守住）；Normal = P50（中位）；Best = P90（冲刺带）。注意与 v0.4 的 Worst=P10 不同——P20 是更灵敏的预警位。"),
    ("与 JY 表手工锚的关系", "JY 表三档（年销 1.2/1.8/3.0 万）在 v0.5 分布中位于 P5/P28/P83 附近——手工锚仍偏保守，作执行接口使用；模型分布作概率读数使用。"),
    ("读数示例", f"「首年 P50 = {pct(y1,50):,.0f} 台」→ 一半的可能性卖得比它多；「P(首年 ≥ 3.0万) = {prob(y1,30000)*100:.1f}%」；「三年累计 P90 = {pct(y3,90):,.0f} 台」→ 20 次里约 2 次三年卖过 {pct(y3,90)/1e4:.0f} 万台。"),
    ("这个分布是什么？", "8,000 次模拟得到的「首年（或三年累计）销量」的全集。每次模拟 = 从先验区间随机抽一组 (m·p·q) 跑一条 Bass 曲线。右尾更长，形状近似对数正态。"),
]
r = 5
for q_, a_ in notes:
    ws3.cell(row=r, column=2, value=q_).font = B_FONT
    ws3.cell(row=r, column=3, value=a_).font = N_FONT
    ws3.cell(row=r, column=3).alignment = WRAP
    ws3.row_dimensions[r].height = 48
    r += 1
ws3.column_dimensions["B"].width = 32
ws3.column_dimensions["C"].width = 100

OUT2 = r"C:\Users\Nan\Documents\GitHub\code_aiken\极印版 aiken\销量概率模型\v0.5\AiKen699销量概率分布_读数表_v05.xlsx"
wb2.save(OUT2)
print("saved:", OUT2)
