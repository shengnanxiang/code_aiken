# -*- coding: utf-8 -*-
"""读取三个 Excel：财务表、JD 站内投放、极印站外投放"""
import sys
sys.stdout.reconfigure(encoding="utf-8")
import openpyxl

def dump(path, out, max_rows=80):
    wb = openpyxl.load_workbook(path, data_only=True)
    lines = [f"### FILE: {path.split(chr(92))[-1]}", f"sheets: {wb.sheetnames}", ""]
    for ws in wb.worksheets:
        lines.append(f"--- SHEET: {ws.title} (rows={ws.max_row}, cols={ws.max_column}) ---")
        n = 0
        for row in ws.iter_rows(values_only=True):
            if all(v is None for v in row):
                continue
            vals = [str(v)[:40] if v is not None else "" for v in row]
            # 截掉尾部空
            while vals and vals[-1] == "":
                vals.pop()
            lines.append(" | ".join(vals))
            n += 1
            if n >= max_rows:
                lines.append(f"... (truncated at {max_rows} rows)")
                break
        lines.append("")
    with open(out, "a", encoding="utf-8") as f:
        f.write("\n".join(lines) + "\n")
    print(f"dumped: {path.split(chr(92))[-1]}")

BASE = r"C:\Users\Nan\Documents\GitHub\code_aiken\极印版 aiken"
OUT = BASE + r"\销量概率模型\_excel_dump.txt"
open(OUT, "w", encoding="utf-8").close()
dump(BASE + r"\AiKen JY Finance_261027.xlsx", OUT)
dump(BASE + r"\销量概率模型\JD26年效果广告分sku_站内.xlsx", OUT)
dump(BASE + r"\销量概率模型\极印投放数据_站外2026YTD1026.xlsx", OUT)
print("all done ->", OUT)
