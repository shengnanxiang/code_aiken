# -*- coding: utf-8 -*-
"""修复 699 读数表的备注列：用模型分位的真实换算，不用 JY 手工锚标签"""
import sys
sys.stdout.reconfigure(encoding="utf-8")
import numpy as np
from openpyxl import load_workbook

XLSX = r"C:\Users\Nan\Documents\GitHub\code_aiken\极印版 aiken\销量概率模型\AiKen699销量概率分布_读数表.xlsx"
wb = load_workbook(XLSX)
ws = wb["销量分位表"]

# 重新计算分位值（与生成脚本一致）
d = np.load(r"C:\Users\Nan\Documents\GitHub\code_aiken\极印版 aiken\销量概率模型\s699_v04_results.npz")
y1, y3 = d["y1"], d["y3"]
pct = lambda a, q_: float(np.percentile(a, q_))

# 找到备注列有内容的行（P10/P50/P90），重写备注
# 表结构：第5行表头，第6行起是 P5...P95，备注在 I 列（第9列）
row = 6
qn = 5
ANCHOR_QN = {10, 50, 90}
while row <= ws.max_row:
    b = ws.cell(row=row, column=2).value
    if isinstance(b, str) and b.startswith("P") and b[1:].isdigit():
        qn = int(b[1:])
        if qn in ANCHOR_QN:
            v1, v3 = pct(y1, qn), pct(y3, qn)
            m1 = v1 / 12.0       # 首年平均月销
            m3 = v3 / 36.0       # 三年平均月销
            label = {10: "Worst 场景锚", 50: "Normal 场景锚", 90: "Best 场景锚"}[qn]
            ws.cell(row=row, column=9,
                    value=f"{label}：首年平均月销 ~{m1:,.0f} 台/月（三年平均 ~{m3:,.0f}）｜ JY 表手工锚：月销1000/1500/2500，年销1.2/1.8/3.0万，对应本分布 P4/P21/P80 附近")
    row += 1

# 同步修正关键概率表里的锚描述（E 列附近区域）——检查下方行
# 关键概率行从分位表结束后开始，labels 引用了 "699 Normal锚" 等，这些是阈值引用没错，无需改

# 修正分布说明 sheet 里最后一条锚描述
ws3 = wb["分布说明"]
for row in range(1, ws3.max_row + 1):
    c = ws3.cell(row=row, column=3)
    if isinstance(c.value, str) and "Worst（月销1,000→年1.2万）" in c.value:
        c.value = c.value.replace(
            "本表口径下：Worst（月销1,000→年1.2万）≈ P10 以内；Normal（月销1,500→年1.8万）≈ P20-P30 区间；Best（月销2,500→年3.0万）≈ P70+",
            "JY 表手工锚（月销1000/1500/2500 → 年销1.2/1.8/3.0万）落在本分布 P4 / P21 / P80 附近；模型 P10/P50/P90 对应首年平均月销 1,213 / 1,844 / 2,733 台"
        )

wb.save(XLSX)
print("fixed & saved")

# 打印验证
print("\n锚位验证：")
for qn in (10, 50, 90):
    v1, v3 = pct(y1, qn), pct(y3, qn)
    print(f"P{qn}: 首年 {v1:,.0f} 台（平均月销 {v1/12:,.0f}）｜ 三年 {v3:,.0f} 台（月均 {v3/36:,.0f}）")

# JY 手工锚在分布中的真实分位
prob_le = lambda a, th: float((a <= th).mean()) * 100
print("\nJY 手工锚的真实分位位置：")
for label, th1 in [("Worst 1.2万", 12_000), ("Normal 1.8万", 18_000), ("Best 3.0万", 30_000)]:
    print(f"{label}: 首年 ≤ 该值的模拟占比 = P{prob_le(y1, th1):.0f}")
