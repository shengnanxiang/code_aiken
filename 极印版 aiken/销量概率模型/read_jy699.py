# -*- coding: utf-8 -*-
"""读取 AiKen JY 699 Finance 表格，理解用户的手工测算结构"""
import sys
sys.stdout.reconfigure(encoding="utf-8")
from openpyxl import load_workbook

PATH = r"C:\Users\Nan\Documents\GitHub\code_aiken\极印版 aiken\AiKen JY 699 Finance_261027.xlsx"
wb = load_workbook(PATH, data_only=True)
print("Sheets:", wb.sheetnames)
for sn in wb.sheetnames:
    ws = wb[sn]
    print(f"\n===== Sheet: {sn} ({ws.max_row}r x {ws.max_column}c) =====")
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 60)):
        vals = []
        for c in row:
            if c.value is not None:
                v = c.value
                if isinstance(v, float):
                    v = round(v, 4)
                vals.append(f"{c.coordinate}:{v}")
        if vals:
            print(" | ".join(vals))
